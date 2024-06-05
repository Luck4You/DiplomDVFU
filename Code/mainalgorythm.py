import workwithdocs
import workwithexcel
import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx import Document
import re

def SumOfZe(dataframe, index_of_disc):
    sum_of_ze = 0
    for index, row in dataframe.iterrows():
        asd = str(dataframe.at[index, 'DISCINDEX'])[:3]
        if index_of_disc[:2] == str(dataframe.at[index, 'DISCINDEX'])[:2]:
            match = re.search(r'\((\d+)\)', str(row['DISCHOURS']))
            if match:
                hours_in_brackets = int(match.group(1))  # Преобразование найденной цифры из строкового типа в целое число
                sum_of_ze += hours_in_brackets
    return sum_of_ze

def MakeOne(df_student,file_save, shablon_path):
    # Открываем документ Word по указанному пути
    doc = Document(shablon_path)

    # Инициализация переменных
    index_Table = 0
    DiscType_now = "1"
    leftTable = None
    rightTable = None
    mainTable = None

    # Создаем пустой DataFrame для хранения информации о курсовых проектах и работах
    KursJob = pd.DataFrame(columns=['TYPECONTROL', 'DISCNAME', 'DISCHOURS', 'OCENKA'])

    # Поиск таблиц в документе Word
    for table in doc.tables:
        if table.cell(0, 0).text == "TABL1TABL":
            leftTable = table
        if table.cell(0, 0).text == "TABL2TABL":
            rightTable = table

    mainTable = leftTable

    # Обработка каждой строки в DataFrame df_student
    for index, row in df_student.iterrows():
        # Проверяем тип контроля, если это курсовой проект или курсовая работа, добавляем в KursJob
        if row['TYPECONTROL'] == 'Курсовой проект' or row['TYPECONTROL'] == 'Курсовая работа':
            new_row = {'TYPECONTROL': row['TYPECONTROL'],
                    'DISCNAME': row['DISCNAME'],
                    'DISCHOURS': row['DISCHOURS'],
                    'OCENKA': row['OCENKA']}
            KursJob.loc[len(KursJob.index)] = new_row
        else:
            match str(row['DISCINDEX'])[:2]:
                case "Б1":
                    discname = row['DISCNAME']
                    if len(discname) >= 56 and index_Table + (len(discname) // 56) <= len(mainTable.rows) - 1:
                        workwithdocs.MergeTable(mainTable, index_Table, (len(discname) // 56))
                        index_Table += (len(discname) // 56)
                    workwithdocs.Insert_Str_from(mainTable,
                                                 index_Table,
                                                 row['DISCNAME'],
                                                 row['DISCHOURS'],
                                                 row['OCENKA'],
                                                 WD_ALIGN_VERTICAL.BOTTOM)
                    index_Table += 1
                case "Б2":
                    

    doc.save(f"C:/Users/pavlov.sa/Desktop/Software/DiplomDVFU-1/Test/Unmerged/{file_save}")

def MakeAll():
    # Указываем пути к файлам
    shablon_path = "C:/Users/pavlov.sa/Desktop/Software/DiplomDVFU-1/Shablons/Diplom.docx"  # Путь к шаблону диплома
    data_file_path = "C:/Users/pavlov.sa/Desktop/Software/DiplomDVFU-1/Test/DataSet.xlsx"  # Путь к файлу с данными

    # Чтение данных из файла Excel в DataFrame
    df = pd.read_excel(data_file_path)

    # Преобразование DataFrame с помощью функции MakeDf из модуля workwithexcel
    df = workwithexcel.MakeDf(df)

    # Получение уникальных значений номеров зачётных книжек
    unique_zachbooks = df['ZACHBOOK'].unique()

    # Создание словаря zk_dfs, содержащего подмножества DataFrame для каждой зачётной книжки
    zk_dfs = {zachbook: df[df['ZACHBOOK'] == zachbook] for zachbook in unique_zachbooks}

    # Для каждой зачётной книжки создаём отдельный файл диплома
    for zachbook, df_part in zk_dfs.items():
        filename = f"{zachbook}.docx"  # Генерируем имя файла на основе номера зачётной книжки
        MakeOne(df_part, filename, shablon_path)  # Вызываем функцию MakeOne для создания файла диплома

    # Записываем отсортированные данные обратно в файл Excel
    df.to_excel("C:/Users/pavlov.sa/Desktop/Software/DiplomDVFU-1/Test/SortedDataSet.xlsx")